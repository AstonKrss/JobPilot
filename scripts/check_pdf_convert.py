"""PDF → DOCX 保版式转换的质量检查。

为什么要有：用户上传的是自己排好版的简历，转换质量直接决定"改完还是不是他那份"。
只靠肉眼看一次不够，需要能重复跑的指标：
    · 文字覆盖率 —— 原 PDF 里的字有多少还留在转换后的 Word 里（丢字就是丢了经历）
    · 结构 —— 转出来是段落还是表格/文本框，DocxEditor 能不能改到
    · 版式锚点 —— 姓名、联系方式、各段标题是否还在原位附近

用法：.venv/Scripts/python.exe scripts/check_pdf_convert.py <简历.pdf>
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from app.rewrite.pdf_convert import convert_pdf_to_docx  # noqa: E402

OUT = ROOT / "data" / "tailored" / "_convert_check.docx"


def pdf_text(path: Path) -> str:
    import pymupdf

    doc = pymupdf.open(str(path))
    try:
        return "".join(page.get_text("text") for page in doc)
    finally:
        doc.close()


def norm(text: str) -> str:
    return "".join(ch for ch in text if not ch.isspace())


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    src = Path(sys.argv[1])
    if not src.exists():
        print(f"文件不存在：{src}")
        return 2

    OUT.parent.mkdir(parents=True, exist_ok=True)
    ok = convert_pdf_to_docx(src, OUT)
    print(f"转换结果: {'成功' if ok else '失败（会回退到重建）'}")
    if not ok:
        return 1

    from docx import Document

    doc = Document(str(OUT))
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    cells = [
        p.text.strip()
        for t in doc.tables
        for row in t.rows
        for cell in row.cells
        for p in cell.paragraphs
        if p.text.strip()
    ]
    docx_text = norm("".join(paras) + "".join(cells))
    original = norm(pdf_text(src))

    # 覆盖率：按 10 字滑窗统计原文字有多少还能在 docx 里找到
    window = 10
    total = max(1, len(original) - window)
    hit = sum(1 for i in range(total) if original[i : i + window] in docx_text)
    coverage = hit / total

    print(f"原 PDF 字数: {len(original)}")
    print(f"Word 字数:   {len(docx_text)}")
    print(f"文字覆盖率:  {coverage:.1%}  （低于 90% 说明转换丢内容）")
    print(f"结构: 正文段落 {len(paras)} 个 / 表格单元 {len(cells)} 个")
    if not cells:
        print("      无表格 —— DocxEditor 可直接改正文段落")
    else:
        print("      含表格 —— DocxEditor 支持 t{r}-c{c}-p{n} 形式的 id")

    # 列出前 12 段，人工确认版式锚点有没有丢
    print("\n前 12 个非空段落:")
    for i, t in enumerate(paras[:12]):
        print(f"  {i:2d}. {t[:70]}")

    return 0 if coverage >= 0.9 else 1


if __name__ == "__main__":
    raise SystemExit(main())
