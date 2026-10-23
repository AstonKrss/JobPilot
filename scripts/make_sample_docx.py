"""生成一份示例 DOCX 简历，用来验证「在原文件上改」这条链路。

刻意做成真实简历的样子：分节标题、公司行加粗、要点用项目符号。
这样能验证编辑后**样式有没有被保留**，而不只是文字变了。
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "samples" / "sample_resume.docx"


def heading(doc, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)


def line(doc, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold


def bullet(doc, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def main() -> int:
    doc = Document()

    title = doc.add_paragraph()
    run = title.add_run("张伟")
    run.bold = True
    run.font.size = Pt(20)

    line(doc, "13800000000 ｜ zhangwei@example.com ｜ 杭州")

    heading(doc, "教育背景")
    line(doc, "清华大学  计算机科学与技术  本科  统招  2015.09 - 2019.06")

    heading(doc, "专业技能")
    line(doc, "Python、Go、MySQL、Redis、Kafka、Docker")

    heading(doc, "工作经历")
    line(doc, "阿里巴巴 · 高级后端工程师 · 杭州    2023.07 - 至今", bold=True)
    bullet(doc, "负责交易链路稳定性，参与同城双活改造。")
    bullet(doc, "推动可观测性建设，故障定位时间缩短。")
    line(doc, "字节跳动 · 后端工程师 · 北京    2019.07 - 2023.06", bold=True)
    bullet(doc, "负责支付核心链路，支撑日均三千万笔交易。")
    bullet(doc, "参与风控引擎重构，规则生效时间从小时级降到分钟级。")

    heading(doc, "项目经历")
    line(doc, "分布式任务调度平台 · 后端负责人", bold=True)
    bullet(doc, "基于 Kubernetes 与 PostgreSQL 实现任务编排，支撑日均五十万任务。")
    bullet(doc, "设计失败重试与幂等机制，任务重复率降到万分之一。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"已生成 {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
