"""给简历"补一个技能词"，不动版式。

背景
    用户看到匹配结果说"缺 Linux"，而他其实会 Linux，只是简历没写全。
    最自然的操作就是"在技能栏加一个词"，而不是重新排版整份简历。

关键约束（这是这个模块存在的理由）
    只改系统画像不够：画像决定了系统怎么推荐，但 HR 看到的是那份 Word/PDF。
    如果画像里有、文件里没有，等于系统在替用户撒谎。
    所以这里在文件里真的追加一个词，而且只追加在技能段落里——
    不去动项目经历、不去编造使用场景。

找不到技能段落就不改，如实返回原因，不做"看起来差不多"的猜测。
"""

from __future__ import annotations

import copy
import logging
from pathlib import Path

from docx import Document

logger = logging.getLogger(__name__)

# 技能栏里常见的分隔符。追加时沿用段落里已有的分隔符，避免出现「、，」混排
SEPARATORS = ("、", "，", ",", "；", ";", " / ", "/")




def _iter_paragraphs(doc):
    """遍历正文段落和表格单元格里的段落。

    PDF 转出来的简历，很多内容在表格里（工作经历、教育经历常是表格）。
    只遍历 doc.paragraphs 会漏掉一半内容。
    """
    for para in doc.paragraphs:
        yield para
    for table in doc.tables:
        seen: set[int] = set()
        for row in table.rows:
            for cell in row.cells:
                key = id(cell._tc)  # 合并单元格会被重复返回
                if key in seen:
                    continue
                seen.add(key)
                for para in cell.paragraphs:
                    yield para


def _skill_line_score(text: str, known_skills: list[str]) -> int:
    """判断这段文本像不像"技能栏"，返回命中数（0 = 不是）。

    踩过的坑：第一版只看"命中多少个已知技能名"，结果选中了项目描述 ——
    「负责开发基于 Spring Boot + Vue 的部门管理系统…」里有 Spring Boot、MySQL、
    Redis 三个技能名，于是把 「、Linux」 追加到了项目描述结尾，
    读起来莫名其妙，而且等于替用户编造"这个项目用了 Linux"。

    所以必须同时满足：
      - 短（技能行通常不超过 80 字）
      - 没有句号（有句号说明是叙述句，不是罗列）
      - 命中至少 2 个已知技能
    """
    if len(text) > 80 or "。" in text:
        return 0
    lower = text.lower()
    hits = sum(1 for s in known_skills if s and len(s) > 1 and s.lower() in lower)
    return hits if hits >= 2 else 0


def append_skill(
    src: Path,
    dst: Path,
    skill: str,
    known_skills: list[str],
) -> dict:
    """在技能段落末尾追加一个词，另存为 dst。

    返回 {"ok": bool, "paragraph": 定位到的段落文本, "reason": 失败原因}
    """
    if not src.exists():
        return {"ok": False, "reason": f"源文件不存在：{src.name}"}
    if src.suffix.lower() != ".docx":
        # PDF 必须先转 Word：在 PDF 上改字会破坏后面的行
        return {"ok": False, "reason": "源文件不是 DOCX，需要先转成 Word"}

    doc = Document(str(src))

    target = None
    best = 0
    for para in _iter_paragraphs(doc):
        text = (para.text or "").strip()
        if not text:
            continue
        score = _skill_line_score(text, known_skills)
        if score > best:
            target, best = para, score

    if target is None:
        return {
            "ok": False,
            "reason": "这份简历没有独立的技能段落（找不到同时含 2 个以上技能名的短行）",
        }

    text = (target.text or "").rstrip()
    if skill.lower() in text.lower():
        return {"ok": False, "reason": f"技能段落里已经有「{skill}」了", "paragraph": text}

    # 沿用段落里已有的分隔符
    sep = next((s for s in SEPARATORS if s in text), "、")

    run = target.add_run(f"{sep}{skill}")
    # 追加的 run 默认用 Normal 样式，会和前面的字不一样。直接复制最后一个 run
    # 的 rPr（字体、字号、颜色、中西文设置都在这），比逐个属性赋值可靠
    if target.runs[:-1]:
        prev = target.runs[-2]
        if prev._element.rPr is not None:
            run._element.insert(0, copy.deepcopy(prev._element.rPr))

    dst.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dst))
    logger.info("技能已追加：%s → %s（段落：%s）", skill, dst.name, text[:40])
    return {"ok": True, "paragraph": text, "separator": sep, "output": dst.name}
