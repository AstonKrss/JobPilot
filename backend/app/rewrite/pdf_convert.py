"""PDF → DOCX 保版式转换。

为什么需要这一层
    用户上传的简历是他自己排的版（很多人觉得自己的模板比任何生成器都好），
    他要的只是「把技能、项目描述改一下」，不是换一套模板重排。
    原来的做法是 PDF 直接按画像重建：内容对了，但版式全变，用户不接受。

    pdf2docx 按 PDF 的实际坐标还原段落、表格、字体、字号、图片，
    转出来的 Word 和原 PDF 基本一致；之后再用 DocxEditor 只改 run 里的文字，
    版式就保住了。

    转换失败或结果明显不可用时返回 False，由调用方回退到重建路径 ——
    宁可给一个版式不同的文件，也不要抛错让用户拿不到东西。
"""

from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger(__name__)

# 转换结果至少要这么多非空段落才算可用。
# 纯图片型 PDF 转出来接近空文档，那种情况重试或回退更有意义
MIN_PARAGRAPHS = 8

# 表格右端留白。实测原简历内容右边距约 24pt（0.85cm），
# 不留出来的话表格会一直顶到裁切线，时间列被切掉半个字
RIGHT_PADDING = int(0.8 * 360000)


def _display_width(text: str) -> int:
    """近似显示宽度：中文按 2，其余按 1。用来按内容分配列宽。"""
    return sum(2 if ord(ch) > 127 else 1 for ch in text)


def _rebalance_single_row(table, usable: int) -> bool:
    """单行表格按内容重分列宽。

    这类表格是从「项目名 + 岗位 + 时间」这种一行三段的排版转来的：
    等比缩放会让长标题列不够宽，标题被折成三行、时间列反而空着。
    按文字宽度分配后，长的那列自然拿到更多位置。
    """
    if len(table.rows) != 1 or len(table.columns) < 2:
        return False
    cells = table.rows[0].cells
    weights = []
    for idx in range(len(table.columns)):
        text = (cells[idx].text or "").strip() if idx < len(cells) else ""
        weights.append(max(_display_width(text), 6))
    total = sum(weights)
    if not total:
        return False
    # 每列至少占 10%，避免时间列被压到换行
    floor = usable * 0.10
    widths = [max(usable * w / total, floor) for w in weights]
    scale = usable / sum(widths)
    for col, width in zip(table.columns, widths):
        col.width = int(width * scale)
    return True


def _align_label_left(table) -> int:
    """把单行表格第一列的长文本改成左对齐。

    pdf2docx 会把「项目名 + 岗位 + 时间」这类行还原成表格后，
    给项目名那一格设成居中 —— 实测原 PDF 里这些项目名的 x 坐标全是 24pt
    （左边距），明明是左对齐。列窄的时候看不出来，列一拉宽就露馅了。
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    if len(table.rows) != 1:
        return 0
    cells = table.rows[0].cells
    if not cells:
        return 0
    first = cells[0]
    if _display_width((first.text or "").strip()) < 10:
        return 0
    changed = 0
    for para in first.paragraphs:
        if para.alignment != WD_ALIGN_PARAGRAPH.LEFT:
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            changed += 1
    return changed


def _widen_email_column(table, table_usable: int) -> int:
    """保证邮箱那一列够宽，别让 @qq.com 断成两行。

    表头区被转成 6 列等宽表格时，邮箱列和其他空列一样宽（约 2.6cm），
    someone@example.com 这类长邮箱放不下就折行，看起来和原版差很多。
    这里把邮箱列拉到够用，其余列按比例让位。
    """
    from docx.shared import Emu

    cols = list(table.columns)
    if len(cols) < 2:
        return 0

    email_idx = None
    for idx in range(len(cols)):
        for row in table.rows:
            cells = row.cells
            if idx < len(cells) and "@" in (cells[idx].text or ""):
                email_idx = idx
                break
        if email_idx is not None:
            break
    if email_idx is None:
        return 0

    widths = [c.width for c in cols]
    if any(w is None for w in widths):
        return 0
    total = sum(widths)
    need = int(3.4 * 360000)  # 17 个英文字符 + 单元格内边距
    if widths[email_idx] >= need or total <= need:
        return 0

    others_total = total - widths[email_idx]
    if others_total <= 0:
        return 0
    remain = table_usable - need
    if remain <= 0:
        return 0
    ratio = remain / others_total
    for idx, col in enumerate(cols):
        col.width = need if idx == email_idx else int(widths[idx] * ratio)
    return 1 if need != widths[email_idx] else 0


def _normalize_layout(path: Path) -> None:
    """修 pdf2docx 的三个通病。

    实测（2026-10-22，两份 2 页中文简历）：
    1. 它把「项目名 + 岗位 + 时间」这一行还原成 1 行 3 列表格，
       列宽总和超过页面可用宽度，项目名被挤成两行；
    2. 行高写的是 exact（固定值），文字一旦换行第二行就直接压在下一行上，
       渲染出来是重叠的黑块；
    3. 表格还带 tblInd 左缩进（实测 480 twips = 0.85cm），
       列宽已经占满可用宽度时，这 0.85cm 会把右端（时间）顶出页面被裁掉。
    这三个都是纯排版参数问题，改完版式就和原 PDF 基本一致了。
    """
    from docx import Document
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    from docx.oxml.ns import qn

    doc = Document(str(path))
    if not doc.sections:
        return
    section = doc.sections[0]
    usable = section.page_width - section.left_margin - section.right_margin
    if not usable:
        return

    changed = 0
    for table in doc.tables:
        # pdf2docx 把页边距设成 0，靠 tblInd 还原原文件的左缩进，
        # 所以不能把 tblInd 归零 —— 那会让表格顶到纸张边缘。
        # 正确做法是：表格可用宽度 = 页面宽 - 左缩进 - 右侧留白
        ind = table._tbl.tblPr.find(qn("w:tblInd"))
        ind_twips = float(ind.get(qn("w:w")) or 0) if ind is not None else 0.0
        ind_emu = int(ind_twips * 635)
        table_usable = usable - ind_emu - RIGHT_PADDING
        if table_usable <= usable * 0.5:
            table_usable = usable - RIGHT_PADDING

        for row in table.rows:
            if row.height_rule == WD_ROW_HEIGHT_RULE.EXACTLY:
                row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
                changed += 1

        if _rebalance_single_row(table, table_usable):
            changed += 1
            changed += _align_label_left(table)
            continue

        # 多行表格（表头区）合并单元格多，按内容分配会算错，等比缩到可用宽度即可
        widths = [col.width for col in table.columns]
        total = sum(w for w in widths if w)
        if total and total > table_usable:
            ratio = table_usable / total
            for col in table.columns:
                if col.width:
                    col.width = int(col.width * ratio)
            changed += 1

        # 等比缩完之后再单独把邮箱列拉够，否则 @qq.com 必然折行
        changed += _widen_email_column(table, table_usable)

    if changed:
        doc.save(str(path))
        logger.info("已修正转换后表格的行高/列宽 %d 处", changed)


def convert_pdf_to_docx(src: Path, dst: Path) -> bool:
    """把 PDF 转成 DOCX。成功且结构可用返回 True。"""
    src, dst = Path(src), Path(dst)
    try:
        from pdf2docx import Converter
    except ImportError:
        logger.warning("未安装 pdf2docx，PDF 无法保版式转换")
        return False

    try:
        converter = Converter(str(src))
        try:
            converter.convert(str(dst))
        finally:
            converter.close()
    except Exception as exc:  # noqa: BLE001 - 转换失败要能回退，不能中断整个定制流程
        logger.warning("PDF 转 DOCX 失败：%s", exc)
        return False

    if not dst.exists() or dst.stat().st_size < 1024:
        logger.warning("PDF 转 DOCX 产物异常：%s", dst)
        return False

    # 结构检查：段落太少说明只转出了图片或空壳
    try:
        from docx import Document

        doc = Document(str(dst))
        texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        table_texts = [
            p.text.strip()
            for t in doc.tables
            for row in t.rows
            for cell in row.cells
            for p in cell.paragraphs
            if p.text.strip()
        ]
        total = len(texts) + len(table_texts)
        if total < MIN_PARAGRAPHS:
            logger.warning("PDF 转 DOCX 后只有 %d 段文字，判定为不可用", total)
            return False
        logger.info("PDF 转 DOCX 完成：%d 段（正文 %d / 表格 %d）", total, len(texts), len(table_texts))
    except Exception as exc:  # noqa: BLE001
        logger.warning("检查转换结果失败：%s", exc)
        return False

    # 文本没问题了，再修排版参数。失败不影响主流程（顶多版式差一点）
    try:
        _normalize_layout(dst)
    except Exception as exc:  # noqa: BLE001
        logger.warning("修正转换版式失败：%s", exc)

    return True
