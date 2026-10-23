"""从结构化画像重建一份干净的 DOCX。

什么时候用
    PDF 是打印结果，没有段落结构，无法"改某一句"。所以 PDF 简历的定制改版
    走重建：用已有的结构化画像（学历/技能/经历/项目）渲染一份单栏 DOCX，
    之后再用同一套 DOCX 编辑工具去改。

    这不是"从零让模型写简历"：内容全部来自画像，模型只改其中几条要点，
    版式由代码固定，出问题一眼能看出来。
"""

from __future__ import annotations

import logging
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

logger = logging.getLogger(__name__)

FONT_CN = "微软雅黑"
FONT_EN = "Calibri"

# 脱敏占位符：重建时保留原位，提醒用户自己补
PLACEHOLDER_HINT = "（原简历中此信息已脱敏，导出后请自行补全）"


def _set_font(run, size: float, bold: bool = False, color: tuple[int, int, int] | None = None) -> None:
    run.font.name = FONT_EN
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    # 中文字体要单独设 eastasia，否则 Word 会用默认宋体
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
    if rfonts is None:
        from docx.oxml.ns import qn

        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    from docx.oxml.ns import qn

    rfonts.set(qn("w:eastAsia"), FONT_CN)


def _heading(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(14)
    para.paragraph_format.space_after = Pt(6)
    run = para.add_run(text)
    _set_font(run, 12, bold=True, color=(48, 86, 211))


def _bullet(doc: Document, text: str) -> None:
    para = doc.add_paragraph(style="List Bullet")
    para.paragraph_format.space_after = Pt(3)
    run = para.add_run(text)
    _set_font(run, 10.5)


def _line(doc: Document, text: str, size: float = 10.5, bold: bool = False, muted: bool = False) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(2)
    run = para.add_run(text)
    _set_font(run, size, bold=bold, color=(110, 115, 122) if muted else None)



def _pdf_header(path: Path) -> tuple[str, str]:
    """从 PDF 首页按字号与坐标取姓名、联系方式。

    不能用 get_text("text") 的顺序：这份 PDF 的文本层是乱的，
    第 0 行是项目名「广东电网…电子签章系统」，第二行是「开发测试员工」，
    按"前 4 行取短文本"会直接把职位当成姓名写进重建的简历里。
    实际排版里姓名是字号最大（22.5pt）且最靠上的那个 span。
    """
    import pymupdf

    doc = pymupdf.open(str(path))
    try:
        if not doc.page_count:
            return "", ""
        spans = []
        for block in doc[0].get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    text = (span.get("text") or "").strip()
                    if not text:
                        continue
                    spans.append(
                        {
                            "size": float(span.get("size") or 0),
                            "y": float(line["bbox"][1]),
                            "x": float(line["bbox"][0]),
                            "t": text,
                        }
                    )
        if not spans:
            return "", ""

        # 姓名：短、不含数字/邮箱特征，字号最大优先，同字号取更靠上的
        cands = [
            s
            for s in spans
            if len(s["t"]) <= 12 and not re.search(r"[0-9@｜|/：:（）()]", s["t"])
        ]
        cands.sort(key=lambda s: (-s["size"], s["y"]))
        name = cands[0]["t"] if cands else ""

        # 联系方式：含手机号或邮箱的那一行，把同一 y 的 span 按 x 拼起来
        # （手机号和邮箱常是两个独立 span）
        anchor_y = None
        for s in spans:
            if re.search(r"1\d{10}|@", s["t"]):
                anchor_y = s["y"]
                break
        contact = ""
        if anchor_y is not None:
            same_line = sorted(
                [s for s in spans if abs(s["y"] - anchor_y) < 2], key=lambda s: s["x"]
            )
            parts = []
            for s in same_line:
                t = s["t"].strip().strip("|｜").strip()
                # 单独的分隔符 span 不要，否则拼出来是「男 | 年龄 | | 手机」
                if t and not re.fullmatch(r"[|｜·、\s]+", t):
                    parts.append(t)
            contact = " · ".join(parts)
        return name, contact
    finally:
        doc.close()


def read_header_text(path: Path) -> tuple[str, str]:
    """从原简历里读姓名与联系方式行。

    为什么从原文件读而不是用画像：画像里的文本是脱敏过的（姓名→[姓名]），
    重建出来会变成占位符。这里只做本地正则/坐标判断，不调用任何模型，
    也不把内容发出去。
    """
    path = Path(path)
    if path.suffix.lower() != ".docx":
        try:
            return _pdf_header(path)
        except Exception as exc:  # noqa: BLE001 - 读不到表头不该阻断主流程
            logger.warning("读取 PDF 表头失败：%s", exc)
            return "", ""

    lines: list[str] = []
    try:
        from docx import Document

        doc = Document(str(path))
        lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    except Exception as exc:  # noqa: BLE001
        logger.warning("读取简历表头失败：%s", exc)
        return "", ""

    name = ""
    contact = ""
    for line in lines[:4]:
        # 姓名：短、无数字、无 @、无邮箱与手机号特征
        if not name and len(line) <= 12 and not re.search(r"[0-9@｜|/]", line):
            name = line
            continue
        if not contact and re.search(r"1\d{10}|@|电话|邮箱", line):
            contact = line
            break
    return name, contact

def build_docx_from_profile(
    profile: dict,
    out_path: Path,
    name: str = "",
    contact: str = "",
) -> Path:
    """按画像渲染单栏简历。返回生成的文件路径。"""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Pt(48)
    section.bottom_margin = Pt(48)
    section.left_margin = Pt(54)
    section.right_margin = Pt(54)

    basic = profile.get("basic") or {}

    # 姓名：直接用原文件里读到的那一行，不经过脱敏
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    _set_font(title.add_run(name or "（请填写姓名）"), 20, bold=True)

    if contact:
        line = doc.add_paragraph()
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        line.paragraph_format.space_after = Pt(10)
        _set_font(line.add_run(contact), 10, color=(110, 115, 122))

    # 教育
    edu_bits = [
        basic.get("school"),
        basic.get("major"),
        basic.get("edu_level"),
        "统招" if basic.get("edu_fulltime") is True else None,
        f"{basic['graduation_year']} 年毕业" if basic.get("graduation_year") else None,
    ]
    edu = " · ".join(str(b) for b in edu_bits if b)
    if edu:
        _heading(doc, "教育背景")
        _line(doc, edu)

    # 技能
    skills = [
        (s if isinstance(s, str) else s.get("name"))
        for s in (profile.get("skills") or [])
    ]
    skills = [s for s in skills if s]
    if skills:
        _heading(doc, "专业技能")
        _line(doc, "、".join(skills))

    # 工作经历
    # 简历的惯例是最近的在最前；抽取出来的顺序不保证，这里显式排序
    exps = sorted(
        profile.get("experiences") or [],
        key=lambda e: (e.get("end") or "9999-99"),
        reverse=True,
    )
    if exps:
        _heading(doc, "工作经历")
        for exp in exps:
            head = " · ".join(
                str(x) for x in (exp.get("company"), exp.get("title"), exp.get("city")) if x
            )
            period = " ~ ".join(
                str(x) for x in (exp.get("start"), exp.get("end") or "至今") if x
            )
            _line(doc, f"{head}    {period}".strip(), bold=True)
            for item in exp.get("highlights") or []:
                if str(item).strip():
                    _bullet(doc, str(item))

    # 项目经历
    projects = profile.get("projects") or []
    if projects:
        _heading(doc, "项目经历")
        for proj in projects:
            head = " · ".join(
                str(x) for x in (proj.get("name"), proj.get("role")) if x
            )
            _line(doc, head, bold=True)
            tech = proj.get("tech") or []
            if tech:
                _line(doc, "技术栈：" + "、".join(str(t) for t in tech), muted=True)
            for item in proj.get("highlights") or []:
                if str(item).strip():
                    _bullet(doc, str(item))

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    logger.info("按画像重建 DOCX：%s", out_path.name)
    return out_path
